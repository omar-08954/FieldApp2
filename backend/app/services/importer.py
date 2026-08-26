"""Fast, fault-isolated Excel import with batched database writes."""
import json
import logging
from difflib import SequenceMatcher
from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session
from sqlalchemy import select

from app.models import ImportBatch, ImportReview, Task, User

logger = logging.getLogger(__name__)
DEFAULT_SUBSCRIPTION = "غير مسجل"
DEFAULT_STATUS = "تم الفحص"
BATCH_SIZE = 250
ALIASES = {
  "technician_name": {"الفني", "اسم الفني", "technician", "employee"},
  "task_number": {"رقم المهمة", "المهمة", "task number", "task", "work order"},
  "subscription_number": {"رقم الاشتراك", "الاشتراك", "subscription", "subscription number"},
  "task_type": {"نوع المهمة", "النوع", "task type", "type"},
  "task_status": {"حالة المهمة", "الحالة", "task status", "status"},
  "city": {"المدينة", "city", "المنطقة"},
  "notes": {"الملاحظات", "notes", "comment"},
}


def _text(value: Any) -> str:
  return "" if value is None else str(value).strip()


def _normalized(value: str) -> str:
  return " ".join(value.replace("ـ", "").split()).casefold()


def _city(value: str) -> str | None:
  value = _text(value)
  if not value:
      return None
  aliases = {"جده": "جدة", "مكه": "مكة", "المدينه": "المدينة", "المدينه المنوره": "المدينة المنورة"}
  return aliases.get(value, value)


def match_technician(name: str, users: list[User]) -> User | None:
  """Match spelling variants and abbreviated Arabic names only when unambiguous."""
  normalized_name = _normalized(name)
  exact = [user for user in users if _normalized(user.full_name) == normalized_name]
  if len(exact) == 1:
      return exact[0]
  source_tokens = set(normalized_name.split())
  candidates = []
  for user in users:
      target = _normalized(user.full_name)
      target_tokens = set(target.split())
      subset = len(source_tokens) >= 2 and source_tokens.issubset(target_tokens)
      similarity = SequenceMatcher(None, normalized_name, target).ratio()
      if subset or similarity >= 0.86:
          candidates.append((subset, similarity, user))
  if len(candidates) == 1:
      return candidates[0][2]
  if candidates:
      candidates.sort(key=lambda item: (item[0], item[1]), reverse=True)
      if len(candidates) == 1 or candidates[0][:2] > candidates[1][:2]:
          return candidates[0][2]
  return None


class ImportValidationError(ValueError):
  def __init__(self, field_name: str, message: str, suggestion: str) -> None:
      super().__init__(message)
      self.field_name = field_name
      self.suggestion = suggestion


def _mapping(headers: list[Any]) -> dict[str, int]:
  normalized = {_text(header).casefold(): index for index, header in enumerate(headers) if _text(header)}
  result: dict[str, int] = {}
  for field, aliases in ALIASES.items():
      for alias in aliases:
          if alias.casefold() in normalized:
              result[field] = normalized[alias.casefold()]
              break
  return result


def _review(
  db: Session,
  batch: ImportBatch,
  row_number: int,
  raw: list[Any],
  exc: Exception,
  action: str,
  values: dict[str, str] | None = None,
) -> None:
  values = values or {}
  field_name = getattr(exc, "field_name", None)
  suggestion = getattr(exc, "suggestion", None)
  db.add(
      ImportReview(
          batch_id=batch.id,
          source_row=row_number,
          task_number=values.get("task_number"),
          technician_name=values.get("technician_name"),
          subscription_number=values.get("subscription_number"),
          task_type=values.get("task_type"),
          task_status=values.get("task_status"), city=values.get("city"), notes=values.get("notes"),
          execution_date=date.fromisoformat(values["execution_date"]) if values.get("execution_date") else None,
          field_name=field_name, suggested_action=suggestion,
          exception_type=type(exc).__name__,
          error_message=str(exc)[:4000],
          postgres_message=str(exc)[:4000] if isinstance(exc, SQLAlchemyError) else None,
          action_taken=action,
          raw_payload=json.dumps({"raw": raw, "values": values}, ensure_ascii=False, default=str),
      )
  )


def task_values(values: dict[str, str]) -> dict[str, Any]:
  """Build a Task payload without ever accepting a spreadsheet primary key."""
  return {
      "technician_id": int(values["_technician_id"]) if values.get("_technician_id") else None,
      "technician_name": values.get("technician_name") or "غير مسجل",
      "task_number": values.get("task_number", ""),
      "subscription_number": values.get("subscription_number") or DEFAULT_SUBSCRIPTION,
      "task_type": values.get("task_type") or "تقني",
      "task_status": values.get("task_status") or DEFAULT_STATUS,
      "city": _city(values.get("city", "")),
      "notes": values.get("notes") or None,
      "execution_date": date.today(),
  }


def _insert_task_batch(
  db: Session,
  batch: ImportBatch,
  rows: list[tuple[int, list[Any], dict[str, str]]],
) -> None:
  """Insert valid rows efficiently and split only failed batches to isolate bad rows."""
  if not rows:
      return
  try:
      with db.begin_nested():
          db.bulk_insert_mappings(Task, [task_values(values) for _, _, values in rows])
  except Exception as exc:
      if len(rows) == 1:
          row_number, raw, values = rows[0]
          logger.warning("Excel row moved to review", extra={"row_number": row_number})
          _review(db, batch, row_number, raw, exc, "moved_to_import_review", values)
          batch.review_rows += 1
          return
      midpoint = len(rows) // 2
      _insert_task_batch(db, batch, rows[:midpoint])
      _insert_task_batch(db, batch, rows[midpoint:])
  else:
      batch.imported_rows += len(rows)


def import_workbook(db: Session, contents: bytes, filename: str, imported_by_id: int | None) -> ImportBatch:
  """Process an Excel workbook in batches while isolating invalid rows."""
  batch = ImportBatch(filename=filename[:255], imported_by_id=imported_by_id)
  db.add(batch)
  db.flush()
  pending: list[tuple[int, list[Any], dict[str, str]]] = []
  technicians = list(db.scalars(select(User).where(User.is_active.is_(True))).all())
  try:
      workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
      try:
          rows = workbook.active.iter_rows(values_only=True)
          headers = list(next(rows, ()))
          mapping = _mapping(headers)
          if "task_number" not in mapping:
              raise ValueError("لم يتم العثور على عمود رقم المهمة في الملف")
          for row_number, raw_tuple in enumerate(rows, start=2):
              raw = list(raw_tuple)
              batch.total_rows += 1
              values: dict[str, str] = {}
              try:
                  values = {
                      field: _text(raw[index]) if index < len(raw) else ""
                      for field, index in mapping.items()
                  }
                  if not values.get("task_number", ""):
                      raise ValueError("رقم المهمة مطلوب")
                  if "technician_name" in mapping:
                      technician_name = values.get("technician_name", "")
                      if not technician_name:
                          raise ImportValidationError("technician_name", "اسم الفني مطلوب", "أدخل اسم الفني أو أضف الفني إلى النظام أولاً")
                      technician = match_technician(technician_name, technicians)
                      if not technician:
                          names = "، ".join(user.full_name for user in technicians[:3])
                          raise ImportValidationError("technician_name", f"الفني «{technician_name}» غير موجود أو غير واضح", f"صحح اسم الفني أو أضفه للنظام. أمثلة أسماء مسجلة: {names}")
                      values["technician_name"] = technician.full_name
                      values["_technician_id"] = str(technician.id)
                  if not values.get("task_type"):
                      raise ImportValidationError("task_type", "نوع المهمة غير موجود", "أدخل نوع المهمة في هذا العمود ثم أعد إدراج الصف")
                  pending.append((row_number, raw, values))
              except Exception as exc:
                  _review(db, batch, row_number, raw, exc, "moved_to_import_review", values)
                  batch.review_rows += 1
              if len(pending) >= BATCH_SIZE:
                  _insert_task_batch(db, batch, pending)
                  pending.clear()
          _insert_task_batch(db, batch, pending)
      finally:
          workbook.close()
      db.commit()
  except Exception as exc:
      logger.exception("Excel file processing failed safely", extra={"import_filename": filename})
      _review(db, batch, 1, [], exc, "file_moved_to_import_review")
      batch.review_rows += 1
      db.commit()
  return batch
